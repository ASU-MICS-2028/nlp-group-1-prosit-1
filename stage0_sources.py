"""Stage 0 follow-up: three more Ewe sources, and what they change.

Our main corpus is one mined web source (the HuggingFace sentence pairs), 35% religious text. Three more sources exist:
    EWE_ENGLISH.csv                  English/Ewe sentence pairs, mostly Jehovah's Witnesses and Bible text (column EWE)
    eweenglishsentence(3).json       a dictionary database export: Glosbe examples and peterlin.pl texts (ee_sentence)
    selected transcribed audios.xlsx University of Ghana Waxal project: transcribed SPOKEN image descriptions
                                     (column Transcription), the only conversational Ewe we have

This script cleans them with the Stage 0 rules plus the extra rules these files need, removes sentences already in
our corpus, splits them with the same hash rule as the main split (so no near-copy crosses train/test), and then asks:
    1. how does our model (trained on the web corpus alone) do on each source's dev set, especially spoken Ewe?
    2. does adding the three sources to training help, and where?
Same model as the main path: modified Kneser-Ney, order 5, whitespace tokens.

Run `uv run python stage0_sources.py` (raw files in data/raw/, see README). Writes results/multisource.json and .md
and data/splits_sources/<source>/{train,dev,test}.ewe.txt. About 10 minutes on a laptop CPU.
"""
import csv
import json
import re
from pathlib import Path

from openpyxl import load_workbook

from stage0_audit import OUT, bucket, clean, fix_eth, is_ewe, nfc
from stage3_smoothing import Smoothed

RAW = Path("data/raw")
SPLITS = Path("data/splits_sources")
RESULTS = Path("results/multisource.json")
MAIN = OUT / "90-5-5"

BINARY = re.compile(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]|\\[0-9A-F]{2}")  # control bytes / escaped binary (broken CSV rows)
MARKUP = re.compile(r"<[^>]+>|https?://\S+|www\.\S+")                     # HTML tags and URLs
ZERO_WIDTH = re.compile(r"[​-‍﻿]")                          # invisible characters that are not whitespace
LOOKALIKES = str.maketrans({"ð": "ɖ", "ε": "ɛ"})                           # lower-case eth and Greek epsilon typed for ɖ and ɛ
LETTER = re.compile(r"[A-Za-zɖƒɣŋɔɛʋƉƑƔŊƆƐƲ]")


def read_sources():
    """Raw Ewe lines from each file, in its own format."""
    with open(RAW / "EWE_ENGLISH.csv", encoding="utf-8", errors="ignore") as f:
        csv_lines = [row["EWE"] for row in csv.DictReader(f) if row.get("EWE")]
    export = json.loads((RAW / "eweenglishsentence(3).json").read_text(encoding="utf-8"))
    table = next(t for t in export if t.get("type") == "table")["data"]  # a PHPMyAdmin export: header, database, table
    dict_lines = [row["ee_sentence"] for row in table if row.get("ee_sentence")]
    sheet = load_workbook(RAW / "selected transcribed audios.xlsx", read_only=True).active
    rows = sheet.iter_rows(values_only=True)
    col = next(rows).index("Transcription")  # the sheet also holds speaker ID, gender, age and device; we use only this
    speech_lines = [str(r[col]) for r in rows if r[col]]
    return {"bible_csv": csv_lines, "dictionary": dict_lines, "speech": speech_lines}


def clean_extra(s):
    """Stage 0 cleaning plus the rules these files need. Returns None for lines to drop."""
    if BINARY.search(s):
        return None
    s = fix_eth(clean(ZERO_WIDTH.sub("", nfc(MARKUP.sub(" ", s))).translate(LOOKALIKES)))
    if len(s.split()) < 2 or not LETTER.search(s):  # at least two words and one letter
        return None
    return s


def religious(lines):
    """Share of sentences naming Yehowa or carrying a chapter:verse reference (a lower bound on Bible/JW text)."""
    return sum(bool(re.search(r"yehowa|\b\d+:\d+\b", s.lower())) for s in lines) / max(len(lines), 1)


def split(lines):
    """The main split's rule: hash of the near-duplicate key, bucket 0-4 test, 5-9 dev, 10-99 train."""
    parts = {"train": [], "dev": [], "test": []}
    for s in lines:
        b = bucket(s)
        parts["test" if b < 5 else "dev" if b < 10 else "train"].append(s)
    return parts


def score(train, devs):
    m = Smoothed(train, 5, "whitespace", method="mkn")
    known = m.models[1].counts[()]
    out = {}
    for name, dev in devs.items():
        pp, bpc, _ = m.evaluate(dev)
        toks = [w for s in dev for w in s.split()]
        out[name] = {"pp": pp, "pp_known": m.evaluate(dev, known_only=True)[0], "bpc": bpc,
                     "oov": sum(w not in known for w in toks) / len(toks)}
    del m
    return out


def main():
    main_train = (MAIN / "train.ewe.txt").read_text(encoding="utf-8").splitlines()
    main_dev = (MAIN / "dev.ewe.txt").read_text(encoding="utf-8").splitlines()
    seen = {s.lower() for p in ("train", "dev", "test") for s in (MAIN / f"{p}.ewe.txt").read_text(encoding="utf-8").splitlines()}

    stats, parts = {}, {}
    for name, raw in read_sources().items():
        cleaned = [c for c in (clean_extra(s) for s in raw) if c]
        unique, dup_main = [], 0
        for s in dict.fromkeys(cleaned):          # dedupe within the source (first occurrence wins)
            if s.lower() in seen:                  # and against everything already in our corpus or an earlier source
                dup_main += 1
                continue
            seen.add(s.lower())
            unique.append(s)
        parts[name] = split(unique)
        folder = SPLITS / name
        folder.mkdir(parents=True, exist_ok=True)
        for p, lines in parts[name].items():
            (folder / f"{p}.ewe.txt").write_text("".join(x + "\n" for x in lines), encoding="utf-8")
        stats[name] = {"raw": len(raw), "kept": len(cleaned), "unique_new": len(unique), "already_in_corpus": dup_main,
                       "train": len(parts[name]["train"]), "dev": len(parts[name]["dev"]), "test": len(parts[name]["test"]),
                       "train_words": sum(len(s.split()) for s in parts[name]["train"]),
                       "religious_share": religious(unique), "passes_letter_filter": sum(map(is_ewe, unique)) / max(len(unique), 1)}
        print(f"{name:<11} raw {len(raw):>7,} | kept {len(cleaned):>7,} | new {len(unique):>7,} | already in corpus {dup_main:>6,} "
              f"| religious {stats[name]['religious_share']:.1%} | passes letter filter {stats[name]['passes_letter_filter']:.1%}", flush=True)
    stats["web (main corpus)"] = {"train": len(main_train), "dev": len(main_dev), "religious_share": religious(main_train),
                                  "train_words": sum(len(s.split()) for s in main_train)}

    devs = {"web": main_dev, **{n: parts[n]["dev"] for n in parts}}
    runs = {"web only (main path)": main_train,
            "web + three sources": main_train + [s for n in parts for s in parts[n]["train"]]}
    scores = {}
    for label, train in runs.items():
        scores[label] = {"train_sentences": len(train), "dev": score(train, devs)}
        print(f"\n{label} ({len(train):,} training sentences)")
        for d, r in scores[label]["dev"].items():
            print(f"   dev = {d:<10} perplexity {r['pp']:>8.1f} | known words {r['pp_known']:>8.1f} | bits/char {r['bpc']:.3f} | unknown words {r['oov']:.1%}", flush=True)

    RESULTS.write_text(json.dumps({"sources": stats, "scores": scores}, indent=1))
    md = "# Three more Ewe sources\n\nModified Kneser-Ney, order 5, whitespace tokens. Each source is split with the main split's hash rule; sentences already in the main corpus are removed.\n\n"
    md += "| source | raw lines | new unique sentences | train / dev / test | religious share | passes the letter filter |\n|---|---:|---:|---|---:|---:|\n"
    for n in parts:
        s = stats[n]
        md += f"| {n} | {s['raw']:,} | {s['unique_new']:,} | {s['train']:,} / {s['dev']:,} / {s['test']:,} | {s['religious_share']:.1%} | {s['passes_letter_filter']:.1%} |\n"
    md += f"| web (main corpus) | | | {stats['web (main corpus)']['train']:,} / {stats['web (main corpus)']['dev']:,} / … | {stats['web (main corpus)']['religious_share']:.1%} | 100% (filtered) |\n"
    md += "\n## Bits per character on each source's dev set (lower is better)\n\n| trained on | " + " | ".join(devs) + " |\n|---|" + "---:|" * len(devs) + "\n"
    for label, r in scores.items():
        md += f"| {label} ({r['train_sentences']:,}) | " + " | ".join(f"{r['dev'][d]['bpc']:.3f}" for d in devs) + " |\n"
    md += "\n## Perplexity on known words, and unknown-word rate\n\n| trained on | " + " | ".join(devs) + " |\n|---|" + "---:|" * len(devs) + "\n"
    for label, r in scores.items():
        md += f"| {label} | " + " | ".join(f"{r['dev'][d]['pp_known']:.1f} ({r['dev'][d]['oov']:.1%})" for d in devs) + " |\n"
    Path("results/multisource.md").write_text(md, encoding="utf-8")


if __name__ == "__main__":
    main()
